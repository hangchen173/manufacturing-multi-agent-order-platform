"""Author fixed manufacturing fixtures; labels never depend on system predictions."""
import argparse
import csv
import hashlib
import json
import sys
from copy import deepcopy
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from infrastructure.vector_store.catalog import load_material_catalog

POLICY = "interview-simulation-v1"
AS_OF = "2026-09-01"


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def make_item(catalog, sku, **changes):
    material = catalog[sku]
    item = {
        "material_name_raw": material["material_name"], "specification_raw": material["specification"],
        "quantity": 20, "unit": material["unit"], "unit_price": material["reference_price"],
        "delivery_date": "2026-09-30", "golden_sku_code": sku,
    }
    item.update(changes)
    return item


def cases(catalog):
    def item(sku, **changes):
        return make_item(catalog, sku, **changes)

    def alias(sku):
        return item(sku, material_name_raw=catalog[sku]["aliases"][0])

    def equivalent(sku):
        return item(sku, specification_raw=catalog[sku]["specification"].replace("×", "*"))

    smoke = [
        ("standard", "auto_approve", [item("FST-001")]),
        ("equivalent_spec", "auto_correct", [equivalent("FST-002")]),
        ("missing_quantity", "manual_review", [item("BRG-001", quantity=None)]),
        ("ambiguous_spec", "manual_review", [item("VLV-001", specification_raw="Q11F-16P DN15", golden_sku_code=None)]),
        ("weight_quantity", "auto_approve", [item("MET-008", quantity=51)]),
        ("high_price", "manual_review", [item("SNS-002", unit_price=100)]),
        ("registered_alias", "auto_correct", [alias("SAF-041")]),
        ("missing_spec", "manual_review", [item("VLV-002", specification_raw=None, golden_sku_code=None)]),
        ("missing_price", "manual_review", [item("TRN-001", unit_price=None)]),
        ("two_rows", "auto_approve", [item("BRG-002"), item("PNE-001")]),
        ("equivalent_spec", "auto_correct", [equivalent("FST-010")]),
        ("past_delivery", "manual_review", [item("SNS-003", delivery_date="2026-08-15")]),
    ]
    acceptance = [
        ("standard", "auto_approve", [item("BRG-010")]),
        ("registered_alias", "auto_correct", [alias("SAF-042")]),
        ("ambiguous_spec", "manual_review", [item("VLV-001", specification_raw="Q11F-16P DN15", golden_sku_code=None)]),
        ("equivalent_spec", "auto_correct", [equivalent("FST-020")]),
        ("missing_quantity", "manual_review", [item("PNE-005", quantity=None)]),
        ("missing_price", "manual_review", [item("TRN-005", unit_price=None)]),
        ("pack_quantity", "manual_review", [item("FST-021", quantity=15)]),
        ("past_delivery", "manual_review", [item("SNS-008", delivery_date="2026-08-20")]),
        ("amount_conflict", "manual_review", [item("BRG-008")]),
        ("registered_alias", "auto_correct", [alias("VLV-020")]),
        ("name_conflict", "manual_review", [item("FST-030", material_name_raw="电磁阀", golden_sku_code=None)]),
        ("weight_quantity", "auto_approve", [item("MET-009", quantity=51)]),
        ("pair_quantity", "auto_approve", [item("SAF-043", quantity=501)]),
        ("tool_quantity", "auto_approve", [item("CUT-022", quantity=21)]),
        ("mixed_rows", "auto_approve", [item("TRN-015"), item("VLV-025"), item("FST-035")]),
        ("registered_alias", "auto_correct", [alias("BRG-015")]),
        ("missing_unit", "auto_approve", [item("FST-031", unit=None)]),
        ("missing_spec", "manual_review", [item("PNE-015", specification_raw=None, golden_sku_code=None)]),
        ("missing_total", "auto_approve", [item("SNS-014")]),
        ("long_30_rows", "auto_approve", [item(f"BRG-{index:03d}") for index in range(1, 31)]),
    ]
    return {"smoke": smoke, "acceptance": acceptance}


def display(value):
    return "" if value is None else str(value)


def source_lines(annotation, independent):
    header = annotation["order_level"]
    lines = ["采购需求单" if independent else "制造业采购订单", "订单编号: " + header["order_number"],
             "客户名称: " + header["customer_name"]]
    for index, item in enumerate(annotation["items"], 1):
        if independent:
            lines += [f"项目 {index:02d} | 品名: {item['material_name_raw']} | 型号: {display(item['specification_raw'])}",
                      f"需求数: {display(item['quantity'])} | 单位: {display(item['unit'])} | 单价: {display(item['unit_price'])} | 交期: {item['delivery_date']}"]
        else:
            lines.append(f"{index} {item['material_name_raw']} | 规格: {display(item['specification_raw'])} | 数量: {display(item['quantity'])} {display(item['unit'])} | 单价: {display(item['unit_price'])} | 交期: {item['delivery_date']}")
    if header["total_amount"] is not None:
        lines.append(f"总金额: {header['total_amount']} 元")
    return lines


def render_document(path, annotation, independent):
    lines = source_lines(annotation, independent)
    if path.suffix == ".txt":
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    elif path.suffix == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "采购明细"
        sheet.append([lines[0]])
        sheet.append(["订单编号", annotation["order_level"]["order_number"]])
        sheet.append(["客户名称", annotation["order_level"]["customer_name"]])
        sheet.append([])
        fields = ["material_name_raw", "specification_raw", "quantity", "unit", "unit_price", "delivery_date"]
        labels = ["品名", "规格型号", "数量", "单位", "单价", "交期"]
        if independent:
            fields = ["material_name_raw", "quantity", "unit", "specification_raw", "unit_price", "delivery_date"]
            labels = ["采购品名", "需求数", "计量单位", "型号参数", "采购单价", "要求交期"]
        sheet.append(["序号"] + labels)
        for index, item in enumerate(annotation["items"], 1):
            sheet.append([index] + [item[key] for key in fields])
        if annotation["order_level"]["total_amount"] is not None:
            sheet.append(["总金额", annotation["order_level"]["total_amount"]])
        for cell in sheet[5]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="245B55")
        for column, width in zip("ABCDEFG", [9, 34, 28, 15, 28, 16, 20]):
            sheet.column_dimensions[column].width = width
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.row_dimensions[row[0].row].height = 28
        workbook.save(path)
    else:
        import fitz
        document = fitz.open()
        page = document.new_page(width=1000, height=max(640, 70 + 25 * len(lines)))
        font = fitz.Font("china-s")
        writer = fitz.TextWriter(page.rect)
        for index, line in enumerate(lines):
            size = 20 if index == 0 else 12
            if font.text_length(line, fontsize=size) > 920:
                raise ValueError(f"Fixture line exceeds page width: {path.name}")
            writer.append((40, 48 + 25 * index), line, font=font, fontsize=size)
        writer.write_text(page)
        if path.suffix == ".pdf":
            document.save(path)
        else:
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(path)
        document.close()


def build(output):
    catalog_path = ROOT / "data/standard_materials.csv"
    catalog = {row["sku_code"]: row for row in load_material_catalog(catalog_path)}
    manifest = []
    for split, entries in cases(catalog).items():
        for index, (scenario, action, items) in enumerate(entries, 1):
            document_id = f"{split}_{index:02d}_{scenario}"
            extensions = ("txt", "xlsx", "pdf", "png") if split == "smoke" else ("png", "pdf", "xlsx", "txt")
            extension = extensions[(index - 1) % 4]
            total = None
            if all(item["quantity"] is not None and item["unit_price"] is not None for item in items):
                total = float(sum(Decimal(str(item["quantity"])) * Decimal(str(item["unit_price"])) for item in items))
            if scenario == "amount_conflict":
                total += 100
            elif scenario == "missing_total":
                total = None
            annotation = {
                "document_id": document_id, "dataset": "interview_fixed_v1", "split": split,
                "policy_version": POLICY, "evaluation_as_of": AS_OF,
                "label_basis": "Fixture source fields and catalog identity, with explicitly chosen policy scenario; not system output",
                "scenario": scenario,
                "order_level": {"order_number": f"INT-{split.upper()}-{index:02d}",
                                "customer_name": "华明制造" if split == "smoke" else "远川机电", "total_amount": total},
                "items": [dict(item, line_index=i) for i, item in enumerate(items, 1)],
                "business_decision": {"action": action, "reason": scenario},
            }
            path = output / split / "orders" / f"{document_id}.{extension}"
            path.parent.mkdir(parents=True, exist_ok=True)
            render_document(path, annotation, split == "acceptance")
            annotation_path = output / split / "annotations" / f"{document_id}.json"
            write_json(annotation_path, annotation)
            manifest.append({"document_id": document_id, "split": split, "file_name": path.name,
                             "document_file": str(path.relative_to(output)), "annotation_file": str(annotation_path.relative_to(output)),
                             "input_sha256": sha256(path), "annotation_sha256": sha256(annotation_path)})
    with (output / "manifest.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(manifest[0]))
        writer.writeheader()
        writer.writerows(manifest)
    write_json(output / "dataset_summary.json", {
        "dataset": "interview_fixed_v1", "policy_version": POLICY, "evaluation_as_of": AS_OF,
        "smoke_count": 12, "acceptance_count": 20, "catalog_sha256": sha256(catalog_path),
        "generator_sha256": sha256(__file__), "model_used_to_generate_labels": False,
        "limitations": ["Synthetic fixtures, not independent customer production data",
                        "Acceptance has separate layouts and was authored before model evaluation; repeated scenarios are intentional",
                        "Packaging multiple 10 applies only to 个/只/件; other units have no invented packaging policy",
                        "No claim of manufacturing domain generalization from 32 cases"],
    })
    return manifest


def validate(output):
    from domain.models import ParsedOrder, OrderItem
    from evaluation.run_evaluation import create_evaluation_orchestrator, validate_material_index
    from evaluation.replay_downstream import replay_order
    from config import Config

    validate_material_index(Config())
    orchestrator = create_evaluation_orchestrator(evaluation_as_of=AS_OF)
    records = []
    for annotation_path in sorted(output.glob("*/annotations/*.json")):
        annotation = json.loads(annotation_path.read_text(encoding="utf-8"))
        parsed = ParsedOrder(**annotation["order_level"], parsing_confidence=.95, items=[
            OrderItem(material_name=item["material_name_raw"], specification=item["specification_raw"],
                      **{key: item[key] for key in ("quantity", "unit", "unit_price", "delivery_date")}, confidence_score=.95)
            for item in annotation["items"]
        ])
        result, _ = replay_order(orchestrator, parsed, annotation["document_id"])
        predicted = [item["sku_code"] for item in result["final_result"]["matched_order"]["items"]]
        expected = [item["golden_sku_code"] for item in annotation["items"]]
        action = result["business_decision"]["action"]
        records.append({"document_id": annotation["document_id"], "passed": predicted == expected and action == annotation["business_decision"]["action"],
                        "expected_action": annotation["business_decision"]["action"], "actual_action": action})
    write_json(output / "downstream_fixture_check.json", {"scope": "fixed golden ParsedOrder replay, not extraction accuracy; zero Qwen calls", "records": records})
    return all(row["passed"] for row in records)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=ROOT / "datasets/self_built/interview_v1")
    parser.add_argument("--validate", action="store_true")
    args = parser.parse_args()
    if (args.output_dir / "manifest.csv").exists():
        parser.error("Output is already frozen; use a new --output-dir to author a new version")
    rows = build(args.output_dir)
    print(f"Authored {len(rows)} fixtures in {args.output_dir}")
    if args.validate and not validate(args.output_dir):
        raise SystemExit("Fixture replay mismatch; inspect the report without rewriting labels to match predictions")


if __name__ == "__main__":
    main()
